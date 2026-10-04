"""Tiered content extraction — implementation-plan.md Step 5.

plain-text -> density heuristic -> OCR. Tier 4 (vision-model fallback) is
deliberately deferred to Step 8, where the pluggable LLM client actually
gets built — a one-off vision call here would just be redone then.
needs_fallback() marks which pages would need it, so the routing logic is
real even though the call itself isn't yet.
"""

import re
from pathlib import Path

from bs4 import BeautifulSoup, Comment
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
import openpyxl
import pdfplumber
import pytesseract
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def needs_fallback(char_count: int, has_image: bool) -> bool:
    """rag-pipeline.md §1's density heuristic, calibrated against 217 real
    pages across 8 files: flag if chars<100, or chars<400 with an image
    present (a page with some text and an image needs a higher bar before
    being trusted as complete, since the image is more likely load-bearing
    when text alone is modest rather than absent)."""
    return char_count < 100 or (char_count < 400 and has_image)


_HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
# Elements whose text is one line of their own.
_LINE_TAGS = {"p", "dt", "dd", "blockquote", "pre", "caption", "figcaption", "address"}
# Anything that means "this element holds blocks, not just inline text".
_BLOCK_TAGS = _HEADING_TAGS | _LINE_TAGS | {"table", "ul", "ol", "li", "div", "section", "article", "header", "footer", "dl"}
# A bold paragraph longer than this is emphasized prose, not a heading.
_MAX_HEADING_CHARS = 100


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def _text(el) -> str:
    return _clean(el.get_text(" "))


def _is_bold_heading(p) -> bool:
    """<p><strong>Grading Algorithm:</strong></p> — how Canvas's editor
    usually marks a section, since most instructors never pick a real
    heading style. Only when the *whole* paragraph is bold: "<strong>Late
    Work:</strong> late work is..." is a sentence with a bold lead-in."""
    text = _text(p)
    if not text or len(text) > _MAX_HEADING_CHARS:
        return False
    bold = " ".join(_text(b) for b in p.find_all(["strong", "b"]))
    return re.sub(r"\s", "", bold) == re.sub(r"\s", "", text)


def extract_html_sections(html: str) -> list[dict]:
    """Canvas HTML (pages, syllabus, assignment descriptions) -> sections in
    page order, [{"heading": str | None, "text": str}].

    A section starts at each heading: <h1>-<h6>, or a paragraph that's bold
    and nothing else. Its text keeps one line per paragraph and list item
    ("- item"), and one line per table row with cells joined by " | " — so a
    grading table stays "12.5% | Project" instead of dissolving into prose.
    A table is read as one unit, headings inside its cells included."""
    if not html.strip():
        return []
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()

    sections = [{"heading": None, "lines": []}]
    inline: list[str] = []

    def add_line(text: str):
        text = _clean(text)
        if text:
            sections[-1]["lines"].append(text)

    def flush_inline():
        add_line(" ".join(inline))
        inline.clear()

    def walk(node):
        for child in node.children:
            name = getattr(child, "name", None)
            if name is None:  # a bare string between elements
                if not isinstance(child, Comment):
                    inline.append(str(child))
                continue
            if name not in _BLOCK_TAGS and not child.find(_BLOCK_TAGS):
                inline.append(child.get_text(" "))  # inline element: part of the running line
                continue
            flush_inline()
            if name in _HEADING_TAGS or (name == "p" and _is_bold_heading(child)):
                sections.append({"heading": _text(child).rstrip(":").strip(), "lines": []})
            elif name == "table":
                for row in child.find_all("tr"):
                    cells = [_text(cell) for cell in row.find_all(["td", "th"], recursive=False)]
                    add_line(" | ".join(c for c in cells if c))
            elif name == "li":
                # A nested list's items get their own lines, after this one.
                nested = [sub.extract() for sub in child.find_all(["ul", "ol"], recursive=False)]
                if _text(child):
                    add_line("- " + _text(child))
                for sub in nested:
                    walk(sub)
            elif name in _LINE_TAGS:
                add_line(_text(child))
            else:
                walk(child)
            flush_inline()

    walk(soup)
    flush_inline()
    return [
        {"heading": s["heading"], "text": "\n".join(s["lines"])}
        for s in sections
        if s["heading"] or s["lines"]
    ]


def extract_html_page(html: str) -> str:
    """Canvas HTML -> plain text: extract_html_sections() joined, a heading
    on its own line above its section's lines."""
    return sections_text(extract_html_sections(html))


def sections_text(sections: list[dict]) -> str:
    parts = []
    for section in sections:
        if section["heading"]:
            parts.append(section["heading"])
        if section["text"]:
            parts.append(section["text"])
    return "\n".join(parts)


def extract_pdf(path: Path) -> list[dict]:
    """One entry per page: text, char_count, has_image, needs_fallback."""
    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            has_image = len(page.images) > 0
            pages.append(
                {
                    "page": i + 1,
                    "text": text,
                    "char_count": len(text),
                    "has_image": has_image,
                    "needs_fallback": needs_fallback(len(text), has_image),
                }
            )
    return pages


def extract_pptx(path: Path) -> list[dict]:
    """One entry per slide: text, char_count, has_image."""
    slides = []
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides):
        text = "".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
        has_image = any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
        slides.append(
            {
                "slide": i + 1,
                "text": text,
                "char_count": len(text),
                "has_image": has_image,
            }
        )
    return slides


def _sections(blocks: list[tuple[str | None, str]]) -> list[dict]:
    """(heading or None, line) pairs -> extract_html_sections()'s shape: a
    new section at each heading, empty ones left out."""
    sections: list[dict] = []
    for heading, line in blocks:
        if heading is not None or not sections:
            sections.append({"heading": heading, "lines": []})
        if line:
            sections[-1]["lines"].append(line)
    return [
        {"heading": s["heading"], "text": "\n".join(s["lines"])}
        for s in sections
        if s["heading"] or s["lines"]
    ]


def extract_docx(path: Path) -> list[dict]:
    """A Word document as sections, split at its headings, in document order
    (paragraphs and tables interleaved). Same shape as extract_html_sections,
    so it's indexed the same way (indexing.sections_to_nodes)."""
    document = docx.Document(path)
    blocks: list[tuple[str | None, str]] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = _clean(block.text)
            style = (block.style.name if block.style is not None else "") or ""
            if text and (style.startswith("Heading") or style == "Title"):
                blocks.append((text, ""))
            elif text:
                blocks.append((None, text))
        elif isinstance(block, Table):
            for row in block.rows:
                # A merged cell is returned once per column it spans.
                cells = list(dict.fromkeys(_clean(c.text) for c in row.cells))
                line = " | ".join(c for c in cells if c)
                if line:
                    blocks.append((None, line))
    return _sections(blocks)


# A data sheet can have tens of thousands of rows; course spreadsheets that
# are worth studying from (a reading list, a schedule) are far smaller.
MAX_SHEET_ROWS = 2000


def _cell_text(value) -> str:
    # Excel stores every number as a float: 8 comes back as 8.0.
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def extract_xlsx(path: Path) -> list[dict]:
    """An Excel workbook as one section per non-empty sheet, each row a
    "a | b | c" line (the same as HTML tables, extract_html_sections)."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sections = []
    try:
        for sheet in book.worksheets:
            lines, skipped = [], 0
            for row in sheet.iter_rows(values_only=True):
                cells = [_clean(_cell_text(v)) for v in row if v is not None and str(v).strip()]
                if not cells:
                    continue
                if len(lines) < MAX_SHEET_ROWS:
                    lines.append(" | ".join(cells))
                else:
                    skipped += 1
            if skipped:
                lines.append(f"[{skipped} more rows not read]")
            if lines:
                sections.append({"heading": sheet.title, "text": "\n".join(lines)})
    finally:
        book.close()
    return sections


def extract_text_file(path: Path) -> list[dict]:
    """A .txt, .md or .csv file as one section. UTF-8, else Latin-1, which
    reads any byte sequence."""
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    text = text.replace("\r\n", "\n").strip()
    return [{"heading": None, "text": text}] if text else []


# Put between a page's own text and whatever OCR adds from its images, so
# the chat model can tell the two apart (OCR output can be noisy).
OCR_MARKER = "[Text read from images on this page]"

_WORD = re.compile(r"[a-z0-9]+")
# Fewer new words than this across a whole page is OCR noise ("Woe", "ate"),
# not a diagram worth adding.
_MIN_NEW_OCR_WORDS = 5


def merge_ocr_text(native: str, ocr: str) -> str:
    """A page's own text plus any OCR lines that add something new.

    OCR used to replace the native text outright whenever needs_fallback()
    fired, which is most slides with a picture on them (25 of 43 pages in
    18-654's course-info deck). On 9 of those, OCR read less than the PDF
    already had: the grading slide lost every label ("12.5% Project" became
    "12.5%"), and "December 2nd" became 'December 2"4'. Keeping the native
    text and adding only OCR lines made mostly of new words keeps what OCR
    is good at (text inside diagrams, calendars, logos) without its misreads
    of text the PDF already had."""
    native, ocr = native.strip(), ocr.strip()
    native_words = set(_WORD.findall(native.lower()))
    if not native_words:
        return ocr or native
    added, new_words = [], set()
    for line in ocr.splitlines():
        words = _WORD.findall(line.lower())
        if not any(len(w) >= 3 and w.isalpha() for w in words):
            continue
        unseen = [w for w in words if w not in native_words]
        if len(unseen) * 2 > len(words):
            added.append(line.strip())
            new_words.update(unseen)
    if len(new_words) < _MIN_NEW_OCR_WORDS:
        return native
    return f"{native}\n\n{OCR_MARKER}\n" + "\n".join(added)


def ocr_pdf_page(path: Path, page_number: int) -> str:
    """Rasterize one PDF page (1-indexed) and run Tesseract on it."""
    with pdfplumber.open(path) as pdf:
        rendered = pdf.pages[page_number - 1].to_image(resolution=200)
        return pytesseract.image_to_string(rendered.original)
