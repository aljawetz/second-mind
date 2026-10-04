"""ingestion.extract_docx / extract_xlsx / extract_text_file — Word, Excel and
plain-text course files (module items in real courses: "Fintech App.docx",
"Sample Table of Selected Papers.xlsx"). Built here with python-docx and
openpyxl, then read back, so no fixture files are needed."""

import docx
import openpyxl

import ingestion


def test_docx_keeps_headings_paragraphs_and_tables_in_order(tmp_path):
    document = docx.Document()
    document.add_paragraph("A budgeting app for students.")
    document.add_heading("Requirements", level=1)
    document.add_paragraph("Track spending by category.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Feature", "Priority"
    table.cell(1, 0).text, table.cell(1, 1).text = "Alerts", "High"
    document.add_heading("Out of scope", level=2)
    document.add_paragraph("Investing.")
    path = tmp_path / "Fintech App.docx"
    document.save(path)

    assert ingestion.extract_docx(path) == [
        {"heading": None, "text": "A budgeting app for students."},
        {"heading": "Requirements", "text": "Track spending by category.\nFeature | Priority\nAlerts | High"},
        {"heading": "Out of scope", "text": "Investing."},
    ]


def test_xlsx_is_one_section_per_sheet_with_rows_as_lines(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Papers"
    sheet.append(["Title", "Year", None])
    sheet.append(["Attention Is All You Need", 2017.0, None])
    sheet.append(["A score", 4.5, None])
    sheet.append([None, None, None])
    book.create_sheet("Empty")
    path = tmp_path / "Sample Table.xlsx"
    book.save(path)

    assert ingestion.extract_xlsx(path) == [
        {"heading": "Papers", "text": "Title | Year\nAttention Is All You Need | 2017\nA score | 4.5"},
    ]


def test_a_huge_sheet_is_cut_off_and_says_so(tmp_path):
    book = openpyxl.Workbook()
    for i in range(ingestion.MAX_SHEET_ROWS + 50):
        book.active.append([f"row {i}"])
    path = tmp_path / "data.xlsx"
    book.save(path)

    [section] = ingestion.extract_xlsx(path)
    lines = section["text"].split("\n")
    assert len(lines) == ingestion.MAX_SHEET_ROWS + 1
    assert lines[-1] == "[50 more rows not read]"


def test_text_files_are_read_whatever_their_encoding(tmp_path):
    utf8 = tmp_path / "notes.md"
    utf8.write_text("# Week 1\nCafé notes", encoding="utf-8")
    latin = tmp_path / "old.txt"
    latin.write_bytes("Café".encode("latin-1"))

    assert ingestion.extract_text_file(utf8) == [{"heading": None, "text": "# Week 1\nCafé notes"}]
    assert ingestion.extract_text_file(latin) == [{"heading": None, "text": "Café"}]
